import os
import json
import pandas as pd
import streamlit as st
from google import genai
from google.genai import types
from openpyxl import load_workbook
from openpyxl.styles import PatternFill

# Configuración de la página de Streamlit
st.set_page_config(
    page_title="Detector de Duplicados en Inventario",
    page_icon="🔍",
    layout="wide"
)

st.title("🔍 Detector inteligente de Duplicados en Excel")
st.markdown("""
Esta aplicación analiza las descripciones de tus materiales utilizando IA, para detectar duplicados.
""")

def normalizar_texto(texto):
    """Realiza una limpieza básica en Python (minúsculas, sin espacios extras)."""
    if pd.isna(texto):
        return ""
    return " ".join(str(texto).lower().strip().split())

# Componente web para subir el archivo Excel
uploaded_file = st.file_uploader("Sube tu archivo de Excel", type=["xlsx", "xls"])

if uploaded_file is not None:
    try:
        # Leer el Excel cargado
        df = pd.read_excel(uploaded_file, sheet_name=0)
        
        st.subheader("Vista previa de los datos originales:")
        st.dataframe(df.head(5), use_container_width=True)
        
        if st.button("🚀 Ejecutar Análisis de Duplicados", type="primary"):
            # Validar y extraer la API Key desde los secretos de Streamlit
            if "GEMINI_API_KEY" not in st.secrets:
                st.error("⚠️ No se encontró la `GEMINI_API_KEY` en la sección de Secrets de Streamlit. Por favor, configúrala en el archivo secrets.toml o en el panel de Streamlit Cloud.")
            else:
                api_key = st.secrets["GEMINI_API_KEY"]
                # Inicializar cliente de Gemini
                client = genai.Client(api_key=api_key)
                
                with st.spinner("Procesando datos y consultando a Gemini... Esto puede tomar unos segundos."):
                    # Validar columnas mínimas esperadas
                    col_codigo = 'código material' if 'código material' in df.columns else df.columns[0]
                    col_desc = 'descripción' if 'descripción' in df.columns else df.columns[1]

                    # Inicializar columnas de resultado
                    df['duplicado'] = 'no'
                    df['código referencia'] = ''
                    df['porcentaje'] = 0
                    df['explicación'] = 'Único / Sin duplicados detectados'

                    # -------------------------------------------------------------
                    # PASO 1: Verificación Simple en Python (Cero Costo / Cero Tokens)
                    # -------------------------------------------------------------
                    df['desc_normalizada'] = df[col_desc].apply(normalizar_texto)
                    df['primer_codigo_grupo'] = df.groupby('desc_normalizada')[col_codigo].transform('first')
                    
                    es_duplicado_exacto = df.duplicated(subset=['desc_normalizada'], keep='first')
                    
                    df.loc[es_duplicado_exacto, 'duplicado'] = 'sí'
                    df.loc[es_duplicado_exacto, 'código referencia'] = df.loc[es_duplicado_exacto, 'primer_codigo_grupo']
                    df.loc[es_duplicado_exacto, 'porcentaje'] = 100
                    df.loc[es_duplicado_exacto, 'explicación'] = 'Coincidencia exacta por normalización de texto (Python)'

                    # -------------------------------------------------------------
                    # PASO 2: Verificación Compleja con la API de Gemini
                    # -------------------------------------------------------------
                    pendientes = df[df['duplicado'] == 'no'].copy()
                    
                    if len(pendientes) > 0:
                        items_a_analizar = pendientes[[col_codigo, col_desc]].to_dict(orient='records')
                        
                        prompt = f"""
                        Eres un experto en inventarios y control de materiales. Analiza la siguiente lista de materiales y detecta si alguno es duplicado o variante de otro elemento de la lista basándote en:
                        - Errores sintácticos o de tipeo (ej. 'vaso plastico' vs 'baso plastico', 'guardabarro' vs 'guardabrrro').
                        - Sinónimos o variaciones de color/términos (ej. 'naraja' vs 'anaranjado', 'rojo' vs 'coloradoo').
                        - Equivalencias en unidades (ej. 'pulgadas' vs '"' o 'pulg').
                        - Cualquier otra similitud semántica clara.

                        IMPORTANTE: Si consideras que el elemento es un duplicado, marca duplicado como "sí", indica el "código referencia" del material original al que está duplicando, un porcentaje de probabilidad y una breve explicación. Si es un elemento único o es el material original principal, deja duplicado en "no", código referencia vacío y porcentaje en 0.

                        Lista en formato JSON:
                        {json.dumps(items_a_analizar, ensure_ascii=False)}

                        Responde estrictamente con un objeto JSON válido que sea una lista de objetos con esta estructura exacta para CADA elemento:
                        [
                          {{
                            "código material": "valor_original",
                            "duplicado": "sí" o "no",
                            "código referencia": "código_del_original (o vacío si es no)",
                            "porcentaje": número entero del 0 al 100,
                            "explicación": "Breve explicación"
                          }},
                          ...
                        ]
                        """

                        try:
                            response = client.models.generate_content(
                                model='gemini-3.7-flash',
                                contents=prompt,
                                config=types.GenerateContentConfig(
                                    automatic_function_calling=types.AutomaticFunctionCallingConfig(disable=True)
                                )
                            )
                            
                            texto_respuesta = response.text.strip()
                            if texto_respuesta.startswith("```json"):
                                texto_respuesta = texto_respuesta[7:-3].strip()
                            elif texto_respuesta.startswith("```"):
                                texto_respuesta = texto_respuesta[3:-3].strip()
                                
                            resultados_gemini = json.loads(texto_respuesta)
                            
                            for res in resultados_gemini:
                                if res.get("duplicado") == "sí":
                                    cod = res.get("código material")
                                    idx = df[df[col_codigo] == cod].index
                                    if len(idx) > 0:
                                        df.loc[idx, 'duplicado'] = 'sí'
                                        df.loc[idx, 'código referencia'] = res.get("código referencia", "")
                                        df.loc[idx, 'porcentaje'] = res.get("porcentaje", 85)
                                        df.loc[idx, 'explicación'] = res.get("explicación", "Detectado por similitud semántica (Gemini)")
                                        
                        except Exception as e:
                            st.error(f"⚠️ Error al procesar con Gemini: {e}")

                    # Limpiar columnas auxiliares
                    df = df.drop(columns=['desc_normalizada', 'primer_codigo_grupo'])

                    # Ordenar el DataFrame agrupando por referencia
                    df['orden_grupo'] = df['código referencia'].apply(lambda x: str(x) if str(x).strip() != '' else 'ZZZ')
                    df = df.sort_values(by=['orden_grupo', col_codigo]).drop(columns=['orden_grupo'])

                    # Guardar temporalmente en disco para aplicar estilos con openpyxl
                    temp_output = "resultado_temp.xlsx"
                    df.to_excel(temp_output, index=False)

                    wb = load_workbook(temp_output)
                    ws = wb.active
                    headers = [cell.value for cell in ws[1]]
                    col_dup_idx = headers.index('duplicado') + 1
                    yellow_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")

                    for row in range(2, ws.max_row + 1):
                        val = ws.cell(row=row, column=col_dup_idx).value
                        if val and str(val).lower() == 'sí':
                            for col in range(1, ws.max_column + 1):
                                ws.cell(row=row, column=col).fill = yellow_fill

                    wb.save(temp_output)
                    
                    # Leer archivo procesado para ofrecerlo en descarga
                    with open(temp_output, "rb") as f:
                        excel_data = f.read()

                    st.success("¡Análisis completado con éxito!")
                    st.subheader("Resultados:")
                    st.dataframe(df, use_container_width=True)

                    # Botón de descarga
                    st.download_button(
                        label="📥 Descargar Excel con Resultados y Formato",
                        data=excel_data,
                        file_name="inventario_analizado.xlsx",
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                    )
    except Exception as e:
        st.error(f"Ocurrió un error al leer el archivo: {e}")